"""
personalize.py — Polza Agency / Тестовое задание
=================================================
Читает xlsx-базу, для каждой компании ищет персональный факт через
Claude API (web_search) и записывает его в столбец «Персонализация».

Использование:
    export ANTHROPIC_API_KEY=sk-ant-...
    pip install anthropic openpyxl
    python personalize.py polza_leads.xlsx
"""

import sys
import time
import anthropic
from openpyxl import load_workbook

# ── Настройки ─────────────────────────────────────────────────────────────
MODEL        = "claude-sonnet-4-20250514"
MAX_TOKENS   = 300
DELAY_SEC    = 1.5   # пауза между запросами (rate-limit)
COL_COMPANY  = 2     # B — название компании
COL_SITE     = 3     # C — сайт
COL_PERSONAL = 8     # H — персонализация
START_ROW    = 2     # первая строка с данными (после заголовка)

PROMPT_TEMPLATE = """
Ты — аутрич-специалист. Найди один конкретный, актуальный факт о компании
«{company}» (сайт: {site}), который можно использовать как персональный крючок
в холодном письме.

Требования к факту:
- 1–2 коротких предложения (до 25 слов суммарно)
- Конкретный и проверяемый (рост, запуск продукта, кейс, награда, экспансия)
- Не рекламный, звучит как наблюдение, а не комплимент
- На русском языке

Выведи ТОЛЬКО сам факт, без пояснений и кавычек.
""".strip()


def get_personalization(client: anthropic.Anthropic, company: str, site: str) -> str:
    """Запрашивает персонализацию у Claude с веб-поиском."""
    response = client.messages.create(
        model=MODEL,
        max_tokens=MAX_TOKENS,
        tools=[{"type": "web_search_20250305", "name": "web_search"}],
        messages=[{
            "role": "user",
            "content": PROMPT_TEMPLATE.format(company=company, site=site)
        }]
    )
    # Достаём текстовый блок из ответа (может быть после tool_use)
    for block in response.content:
        if block.type == "text" and block.text.strip():
            return block.text.strip()
    return ""


def main():
    if len(sys.argv) < 2:
        print("Использование: python personalize.py <файл.xlsx>")
        sys.exit(1)

    path = sys.argv[1]
    wb   = load_workbook(path)
    ws   = wb.active
    client = anthropic.Anthropic()   # читает ANTHROPIC_API_KEY из env

    total   = ws.max_row - 1          # строк с данными
    updated = 0

    print(f"Обрабатываем {total} компаний из '{path}'...")

    for row in ws.iter_rows(min_row=START_ROW, max_row=ws.max_row):
        company_cell = row[COL_COMPANY - 1]
        site_cell    = row[COL_SITE - 1]
        personal_cell = row[COL_PERSONAL - 1]

        company = str(company_cell.value or "").strip()
        site    = str(site_cell.value or "").strip()

        if not company or personal_cell.value:
            # Пропускаем пустые строки и уже заполненные
            continue

        print(f"  [{updated + 1}/{total}] {company} ...", end=" ", flush=True)

        try:
            fact = get_personalization(client, company, site)
            personal_cell.value = fact
            updated += 1
            print("✓")
        except Exception as exc:
            print(f"ОШИБКА: {exc}")

        time.sleep(DELAY_SEC)

    wb.save(path)
    print(f"\nГотово. Заполнено {updated} строк. Файл сохранён: {path}")


if __name__ == "__main__":
    main()
