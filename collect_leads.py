"""
collect_leads.py — Polza Agency / Тестовое задание
===================================================
Собирает новые B2B-компании из России через Claude API (web_search)
и дописывает их в xlsx-базу.

Использование:
    export ANTHROPIC_API_KEY=sk-ant-...
    pip install anthropic openpyxl
    python collect_leads.py [--file polza_leads.xlsx] [--segment SaaS] [--count 20]

Аргументы:
    --file      путь к xlsx (по умолчанию polza_leads.xlsx)
    --segment   сегмент для поиска (SaaS, логистика, HR, консалтинг ...)
    --count     сколько компаний найти (по умолчанию 10)
"""

import argparse
import json
import time
import anthropic
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment

MODEL      = "claude-sonnet-4-20250514"
MAX_TOKENS = 2000
DELAY_SEC  = 2.0

SEARCH_PROMPT = """
Найди {count} реальных российских B2B-компаний в сегменте «{segment}».
Критерии: есть свой продукт или услуга для бизнеса, активно ищут клиентов,
есть отдел продаж. Небольшие и средние компании предпочтительнее крупных.

Верни ТОЛЬКО JSON-массив без пояснений, строго в формате:
[
  {{
    "company": "Название компании",
    "site": "домен.ru",
    "contact_name": "Имя Фамилия (реальный ЛПР если известен, иначе должность)",
    "email": "почта@домен.ru",
    "product": "Краткое описание продукта/услуги",
    "segment": "{segment}"
  }},
  ...
]
Только JSON, никакого markdown.
""".strip()


def fetch_leads(client: anthropic.Anthropic, segment: str, count: int) -> list[dict]:
    """Запрашивает список компаний у Claude с веб-поиском."""
    response = client.messages.create(
        model=MODEL,
        max_tokens=MAX_TOKENS,
        tools=[{"type": "web_search_20250305", "name": "web_search"}],
        messages=[{
            "role": "user",
            "content": SEARCH_PROMPT.format(count=count, segment=segment)
        }]
    )
    # Ищем JSON в текстовом блоке
    for block in response.content:
        if block.type == "text":
            text = block.text.strip()
            # Убираем возможные markdown-обёртки
            if text.startswith("```"):
                text = text.split("```")[1]
                if text.startswith("json"):
                    text = text[4:]
            try:
                data = json.loads(text)
                if isinstance(data, list):
                    return data
            except json.JSONDecodeError:
                pass
    return []


def append_to_xlsx(path: str, leads: list[dict]):
    """Дописывает новые строки в существующий xlsx (или создаёт новый)."""
    try:
        wb = load_workbook(path)
        ws = wb.active
        next_row = ws.max_row + 1
        # Определяем следующий порядковый номер
        last_num = ws.cell(ws.max_row, 1).value
        counter  = (last_num or 0) + 1
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.title = "База лидов"
        headers = ["#", "Компания", "Сайт", "ЛПР (имя)", "Email",
                   "Продукт/услуга", "Сегмент", "Персонализация", "Статус"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, name="Arial")
        next_row = 2
        counter  = 1

    font  = Font(name="Arial", size=10)
    align = Alignment(wrap_text=True, vertical="top")

    for lead in leads:
        row_data = [
            counter,
            lead.get("company", ""),
            lead.get("site", ""),
            lead.get("contact_name", ""),
            lead.get("email", ""),
            lead.get("product", ""),
            lead.get("segment", ""),
            "",        # персонализация — позже через personalize.py
            "Новый",
        ]
        ws.append(row_data)
        for cell in ws[ws.max_row]:
            cell.font      = font
            cell.alignment = align
        counter += 1

    wb.save(path)
    print(f"Добавлено {len(leads)} строк в '{path}'")


def main():
    parser = argparse.ArgumentParser(description="Сбор лидов для Polza Agency")
    parser.add_argument("--file",    default="polza_leads.xlsx", help="Путь к xlsx")
    parser.add_argument("--segment", default="B2B SaaS Россия", help="Сегмент поиска")
    parser.add_argument("--count",   type=int, default=10,       help="Кол-во компаний")
    args = parser.parse_args()

    client = anthropic.Anthropic()

    print(f"Ищем {args.count} компаний в сегменте «{args.segment}»...")
    leads = fetch_leads(client, args.segment, args.count)
    print(f"Найдено: {len(leads)}")

    if leads:
        append_to_xlsx(args.file, leads)
        print("Готово! Запустите personalize.py для добавления персонализации.")
    else:
        print("Компании не найдены — проверьте API-ключ и сегмент.")


if __name__ == "__main__":
    main()
